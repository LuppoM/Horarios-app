from datetime import date, timedelta
import os
import re

import tkinter as tk
from tkinter import Toplevel, filedialog, messagebox, ttk

from abmmodmoi import mostrar_abmmodmoi
from modmoimodelo import ModMoiModel
from modmoivista import ModMoiView
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ModMoiController:

    def __init__(self, parent, conexion):
        self.modelo = ModMoiModel(conexion)
        self.vista = ModMoiView(parent, self)
        self._configurar_botones()
        self._configurar_filtro()

        # Evento al cambiar entre Semana y Fin de Semana
        self.vista.combo_vista.bind(
            "<<ComboboxSelected>>", lambda e: self.cargar_datos()
        )

        self.cargar_datos()

    # ------------------- CARGA DE DATOS -------------------
    def cargar_datos(self, filtro_columna=None, filtro_valor=None):
        try:
            # Evaluar cuál tabla se seleccionó en el Combobox
            tipo_vista = self.vista.combo_vista.get()

            if tipo_vista == "Fin de Semana":
                columnas, datos = self.modelo.obtener_datos_finde(
                    filtro_columna, filtro_valor
                )
            else:
                columnas, datos = self.modelo.obtener_datos(
                    filtro_columna, filtro_valor
                )

            self.vista.columnas = columnas
            self.vista.datos_originales = datos
            self._actualizar_tabla()
            self.contar_mano_de_obra()
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def _actualizar_tabla(self):
        tree = self.vista.tree
        tree.delete(*tree.get_children())
        tree["columns"] = self.vista.columnas

        # Configurar encabezados y formato
        for col in self.vista.columnas:
            tree.heading(col, text=col)
            tree.column(col, anchor="center", width=150)

        # Insertar filas
        for fila in self.vista.datos_originales:
            tree.insert("", "end", values=fila)

        # Actualizar combobox de columnas para el filtro según la tabla visible
        self.vista.combo_columnas["values"] = self.vista.columnas
        if self.vista.columnas:
            self.vista.combo_columnas.current(0)

    def contar_mano_de_obra(self):
        try:
            tipo_vista = self.vista.combo_vista.get()
            es_finde = tipo_vista == "Fin de Semana"

            mod_count, moi_count = self.modelo.contar_mano_de_obra(
                es_finde=es_finde
            )
            self.vista.contador_label.config(
                text=(
                    f"MOD (Directa): {mod_count}   |   MOI (Indirecta):"
                    f" {moi_count}"
                )
            )
        except Exception:
            self.vista.contador_label.config(
                text="Error al contar mano de obra"
            )

    # ------------------- BOTONES -------------------
    def _configurar_botones(self):
        v = self.vista
        v.crear_boton(
            "Agregar",
            lambda: self._pedir_contrasena(
                lambda: mostrar_abmmodmoi(
                    v, self.modelo.conexion, self.cargar_datos
                )
            ),
            10,
        )
        v.crear_boton(
            "Modificar",
            lambda: self._pedir_contrasena(self.modificar_ot),
            140,
        )
        v.crear_boton(
            "Eliminar", lambda: self._pedir_contrasena(self.eliminar_ot), 270
        )
        v.crear_boton("Exportar a Excel", self.exportar_bd_a_excel, 400)
        v.crear_boton("Cerrar", v.destroy, 530)

    # ------------------- FILTRO -------------------
    def _configurar_filtro(self):
        v = self.vista
        v.entrada_filtro.bind("<Return>", lambda event: self.aplicar_filtro())
        if "children" in v.filtro_frame.__dict__:
            for child in v.filtro_frame.children.values():
                if isinstance(child, ttk.Button):
                    child.config(command=self.aplicar_filtro)

    def aplicar_filtro(self):
        v = self.vista
        columna = v.combo_columnas.get()
        texto = v.entrada_filtro.get().strip()

        if not columna:
            messagebox.showwarning("Filtro", "Por favor selecciona una columna.")
            return

        if texto == "":
            self.cargar_datos()
            return

        if not re.match(r"^[a-zA-Z0-9\s:áÁ-]+$", texto):
            messagebox.showerror(
                "Filtro inválido",
                "El filtro contiene caracteres no permitidos.\nUsa solo letras,"
                " números, espacios, ':' y '-'",
            )
            return

        self.cargar_datos(filtro_columna=columna, filtro_valor=texto)

    # ------------------- MODIFICAR / ELIMINAR -------------------
    def modificar_ot(self):
        tree = self.vista.tree
        seleccionado = tree.focus()
        if not seleccionado:
            messagebox.showwarning(
                "Modificar OT",
                "Por favor selecciona una fila para modificar.",
            )
            return
        datos_fila = tree.item(seleccionado, "values")
        if not datos_fila:
            return
        datos_dict = dict(zip(self.vista.columnas, datos_fila))
        mostrar_abmmodmoi(
            self.vista, self.modelo.conexion, self.cargar_datos, datos_dict
        )

    def eliminar_ot(self):
        tree = self.vista.tree
        seleccionado = tree.focus()
        if not seleccionado:
            messagebox.showwarning(
                "Eliminar OT", "Por favor selecciona una fila para eliminar."
            )
            return
        datos_fila = tree.item(seleccionado, "values")
        if not datos_fila:
            return
        respuesta = messagebox.askyesno(
            "Eliminar OT",
            f"¿Seguro que deseas eliminar el operario con ID {datos_fila[0]}?",
        )
        if not respuesta:
            return
        try:
            self.modelo.eliminar_ot(datos_fila[0])
            messagebox.showinfo("Éxito", "El operario ha sido eliminado.")
            self.cargar_datos()
        except Exception as e:
            messagebox.showerror("Error", str(e))

    # ------------------- CONTRASEÑA -------------------
    def _pedir_contrasena(self, callback):
        v = self.vista
        if v.acceso_autorizado:
            callback()
            return

        ventana_contra = Toplevel(v)
        ventana_contra.title("Ingrese contraseña")
        ventana_contra.geometry("300x150")
        ventana_contra.resizable(False, False)
        ventana_contra.transient(v)
        ventana_contra.grab_set()

        tk.Label(ventana_contra, text="Contraseña:", font=("Segoe UI", 11)).pack(
            pady=10
        )
        entrada = ttk.Entry(ventana_contra, show="*")
        entrada.pack()
        entrada.focus()

        def validar():
            if entrada.get() == "Mario456":
                v.acceso_autorizado = True
                ventana_contra.destroy()
                callback()
            else:
                messagebox.showerror("Acceso denegado", "Contraseña incorrecta")

        ttk.Button(ventana_contra, text="Aceptar", command=validar).pack(pady=10)
        ventana_contra.bind("<Return>", lambda e: validar())

    # ------------------- EXPORTAR A EXCEL -------------------
    def exportar_bd_a_excel(self):
        try:
            # 1. Detectar tabla seleccionada
            tipo_vista = self.vista.combo_vista.get()
            es_finde = tipo_vista == "Fin de Semana"

            # 2. Calcular número de semana para el nombre
            fecha_proxima = date.today() + timedelta(days=7)
            semana_actual = self.modelo.obtener_numero_semana_personalizado(
                fecha_proxima
            )

            # 3. Obtener datos y configurar nombres según tabla activa
            if es_finde:
                columnas, datos = self.modelo.obtener_datos_finde()
                nombre_hoja = "Horarios Finde"
                nombre_archivo_default = (
                    f"Consolidado Finde semana {semana_actual}.xlsx"
                )
            else:
                columnas, datos = self.modelo.obtener_datos()
                nombre_hoja = "Horarios Semana"
                nombre_archivo_default = (
                    f"Consolidado Semana {semana_actual}.xlsx"
                )

            if not datos:
                messagebox.showinfo(
                    "Exportar",
                    f"No hay datos registrados en '{nombre_hoja}' para exportar.",
                )
                return

            ruta_archivo = filedialog.asksaveasfilename(
                title=f"Guardar Excel ({nombre_hoja})",
                defaultextension=".xlsx",
                filetypes=[("Archivos Excel", "*.xlsx")],
                initialfile=nombre_archivo_default,
            )

            if not ruta_archivo:
                return

            # 4. Generar el libro Excel
            wb = Workbook()
            ws = wb.active
            ws.title = nombre_hoja

            encabezado_font = Font(bold=True, color="FFFFFF")
            encabezado_fill = PatternFill("solid", fgColor="4F81BD")
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            # Encabezados
            for i, col in enumerate(columnas, start=1):
                celda = ws.cell(row=1, column=i, value=col)
                celda.font = encabezado_font
                celda.fill = encabezado_fill
                celda.border = thin_border
                celda.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            # Filas de datos
            for fila_idx, fila in enumerate(datos, start=2):
                for col_idx, valor in enumerate(fila, start=1):
                    celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
                    celda.border = thin_border
                    celda.alignment = Alignment(
                        vertical="center", wrap_text=True
                    )

            # Ajuste de ancho de columnas
            for col_idx, col in enumerate(columnas, start=1):
                max_length = len(str(col))
                for fila in datos:
                    val_celda = fila[col_idx - 1]
                    texto_val = str(val_celda) if val_celda is not None else ""
                    if len(texto_val) > max_length:
                        max_length = len(texto_val)

                letra_col = get_column_letter(col_idx)
                ws.column_dimensions[letra_col].width = max(
                    15, min(max_length + 2, 50)
                )

            wb.save(ruta_archivo)

            messagebox.showinfo(
                "Exportación exitosa",
                f"Consolidado de {nombre_hoja} exportado correctamente a:\n{ruta_archivo}",
            )

        except Exception as e:
            messagebox.showerror(
                "Error al exportar", f"No se pudo generar el archivo Excel:\n{e}"
            )